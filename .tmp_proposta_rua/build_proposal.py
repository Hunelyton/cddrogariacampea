from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "proposta_comercial_inventario_rua_drogarias_campea.docx"
LOGO = ROOT / "src" / "assets" / "logo-drogaria-campea.png"

GREEN = "176B45"
LIGHT_GREEN = "EAF4EF"
LIGHT_GRAY = "F4F5F6"
BORDER = "D9D9D9"
TEXT = RGBColor(31, 41, 55)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(6)

for style_name, size in (("Heading 1", 15), ("Heading 2", 11.5)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True

header = section.header
header.distance = Cm(0.7)
header_p = header.paragraphs[0]
header_p.text = "PROPOSTA COMERCIAL   |   INVENTÁRIO DE RUA"
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_p.runs[0].font.name = "Aptos"
header_p.runs[0].font.size = Pt(8)
header_p.runs[0].font.bold = True
header_p.runs[0].font.color.rgb = RGBColor(80, 80, 80)

if LOGO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Cm(4.2))

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Proposta comercial para melhorias no sistema de inventário de rua")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Drogarias Campeã")
run.bold = True
run.font.size = Pt(13)
subtitle.paragraph_format.space_after = Pt(16)

summary = doc.add_table(rows=4, cols=2)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
summary.columns[0].width = Cm(4.2)
summary.columns[1].width = Cm(11.8)
summary_data = [
    ("Proponente", "Hunelyton Mendes Lima | CPF 372.044.388-41"),
    ("Contratante", "Drogarias Campeã | CNPJ 46.756.296/0001-76"),
    ("Objeto", "Inclusão do módulo de inventário de rua e melhorias relacionadas"),
    ("Investimento", "R$ 750,00 (setecentos e cinquenta reais)"),
]
for index, (label, value) in enumerate(summary_data):
    left, right = summary.rows[index].cells
    left.text = label
    right.text = value
    set_cell_shading(left, GREEN)
    for run in left.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    if index == 3:
        set_cell_shading(right, LIGHT_GREEN)
        right.paragraphs[0].runs[0].font.bold = True
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_borders(cell)

doc.add_paragraph()

doc.add_heading("1 Objetivo da proposta", level=1)
doc.add_paragraph(
    "Esta proposta apresenta as condições para desenvolver, implantar e validar uma nova opção de inventário de rua no sistema utilizado pela Drogarias Campeã. A melhoria manterá o inventário de picking existente e acrescentará um módulo independente, permitindo que a equipe selecione o tipo de inventário antes de iniciar o trabalho."
)
doc.add_paragraph(
    "O novo módulo organizará o cadastro, a importação das contagens, o confronto dos dados, os ajustes operacionais e a emissão de relatórios específicos para o inventário de rua. O objetivo é oferecer uma operação clara, rastreável e separada do picking, sem alterar os dados já utilizados no processo atual."
)

doc.add_heading("2 Escopo dos serviços", level=1)

items = [
    ("2.1 Seleção do tipo de inventário", "Criação de uma tela inicial para escolha entre os módulos Picking e Rua. Cada opção direcionará o usuário para seu respectivo painel e utilizará armazenamento separado, evitando mistura entre cadastros e contagens."),
    ("2.2 Cadastro do inventário de rua", "Importação do arquivo de cadastro nos formatos XLSX, XLS ou TXT, respeitando a ordem definida para as colunas Crachá, Inventário Escopo, Número da Contagem, Tipo de Coleta, Código Localizador, Código e Quantidade. Todos os campos serão preservados e exibidos no sistema."),
    ("2.3 Importação da contagem", "Importação do arquivo TXT gerado pelo coletor, utilizando o mesmo mapeamento de sete colunas. As informações importadas ficarão disponíveis para consulta, pesquisa e confronto com o cadastro do inventário de rua."),
    ("2.4 Confronto e identificação de divergências", "Comparação dos dados pelo Código e pela Quantidade. Quando um código aparecer em mais de uma linha, as quantidades serão consolidadas para o confronto. O sistema identificará itens conferidos, divergentes, não contados e não cadastrados, além de indicar faltas e sobras em unidades."),
    ("2.5 Contagem manual e ajustes", "Inclusão manual de registros na aba Contagem e disponibilização do campo Quantidade Ajustada. A quantidade original será preservada, e o confronto, os indicadores e os relatórios considerarão a quantidade ajustada quando houver alteração."),
    ("2.6 Pesquisa e consulta", "Disponibilização de pesquisa em todos os campos das abas Cadastro, Contagem e Divergências. A consulta poderá localizar também códigos já conferidos, permitindo verificar registros mesmo quando não houver diferença."),
    ("2.7 Indicadores do inventário de rua", "Apresentação de métricas de produtos cadastrados, itens e unidades contadas, localizadores, divergências ativas, faltas, sobras, erros por operador, produtos não cadastrados, inclusões manuais e ajustes realizados."),
    ("2.8 Exportações e relatório", "Exportação da contagem nos formatos TXT e XLSX. Inclusão de relatório completo em PDF com resumo geral, resumo operacional, análise por operador, produtos não cadastrados, maiores faltas e sobras, divergências, cadastro completo e contagem detalhada com quantidade original, quantidade ajustada e origem do registro."),
    ("2.9 Limpeza independente dos dados", "Inclusão de uma opção para limpar os dados do inventário de rua sem remover ou modificar as informações armazenadas no módulo de picking."),
    ("2.10 Testes e homologação", "Disponibilização da atualização para testes com a equipe da Drogarias Campeã. Serão realizados ajustes pontuais diretamente relacionados às funções descritas nesta proposta, antes da liberação final."),
]
for heading, body in items:
    doc.add_heading(heading, level=2)
    doc.add_paragraph(body)

doc.add_heading("3 Fluxo de funcionamento", level=1)
flow = [
    "O usuário acessa o sistema e escolhe entre Inventário Picking e Inventário Rua.",
    "No módulo Rua, importa o cadastro no layout acordado.",
    "Em seguida, importa a contagem TXT do coletor ou inclui registros manualmente.",
    "O sistema consolida os códigos e confronta as quantidades cadastradas e contadas.",
    "A equipe consulta as divergências, pesquisa qualquer campo e realiza ajustes quando necessário.",
    "Após a conferência, a contagem pode ser exportada em TXT ou XLSX e o relatório completo pode ser emitido em PDF.",
]
for index, text in enumerate(flow, 1):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)

doc.add_heading("4 Entregáveis", level=1)
deliverables = [
    "Tela inicial de seleção entre Picking e Rua.",
    "Módulo de inventário de rua com dados independentes.",
    "Importação do cadastro e da contagem conforme o layout definido.",
    "Confronto por código e quantidade, com indicadores operacionais.",
    "Inclusão manual e ajuste das quantidades contadas.",
    "Pesquisa completa nas abas do módulo.",
    "Exportação da contagem em TXT e XLSX.",
    "Relatório completo do inventário de rua em PDF.",
    "Versão para testes e ajustes de homologação dentro do escopo.",
]
for text in deliverables:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("5 Prazo de entrega", level=1)
doc.add_paragraph(
    "O prazo para disponibilização da atualização será de até 10 dias úteis, contados a partir da confirmação do pagamento e do recebimento de todos os arquivos e informações necessários para validação do layout. A entrega será realizada inicialmente em ambiente de testes para homologação da equipe responsável pelo inventário."
)
doc.add_paragraph(
    "O prazo poderá ser revisto caso sejam solicitadas alterações fora do escopo, haja mudança no layout dos arquivos, indisponibilidade de acesso ao sistema ou dependência de validação por parte da contratante."
)

doc.add_heading("6 Valor da proposta", level=1)
value_table = doc.add_table(rows=2, cols=2)
value_table.alignment = WD_TABLE_ALIGNMENT.CENTER
value_table.autofit = False
value_table.columns[0].width = Cm(10.8)
value_table.columns[1].width = Cm(5.2)
value_table.cell(0, 0).text = "Serviço"
value_table.cell(0, 1).text = "Valor"
value_table.cell(1, 0).text = "Desenvolvimento e implantação do módulo de inventário de rua"
value_table.cell(1, 1).text = "R$ 750,00"
set_repeat_table_header(value_table.rows[0])
for c in value_table.rows[0].cells:
    set_cell_shading(c, GREEN)
    for run in c.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
for row in value_table.rows:
    for cell in row.cells:
        set_cell_margins(cell, 140, 150, 140, 150)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
value_table.cell(1, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
value_table.cell(1, 1).paragraphs[0].runs[0].font.bold = True
value_table.cell(1, 1).paragraphs[0].runs[0].font.size = Pt(12)
set_cell_shading(value_table.cell(1, 1), LIGHT_GREEN)
doc.add_paragraph("Valor total: R$ 750,00 (setecentos e cinquenta reais). O valor contempla o desenvolvimento, a implantação, os testes e os ajustes diretamente relacionados ao escopo descrito nesta proposta.")

doc.add_heading("7 Forma de pagamento", level=1)
doc.add_paragraph("O pagamento deverá ser realizado via PIX, conforme os dados abaixo:")
payment = doc.add_table(rows=2, cols=2)
payment.alignment = WD_TABLE_ALIGNMENT.LEFT
payment_data = [("Chave PIX", "372.044.388-41"), ("Favorecido", "Hunelyton Mendes Lima")]
for i, (label, value) in enumerate(payment_data):
    payment.cell(i, 0).text = label
    payment.cell(i, 1).text = value
    set_cell_shading(payment.cell(i, 0), LIGHT_GRAY)
    payment.cell(i, 0).paragraphs[0].runs[0].font.bold = True
    for cell in payment.rows[i].cells:
        set_cell_margins(cell)
        set_cell_borders(cell)
doc.add_paragraph("A confirmação do pagamento será considerada o aceite formal desta proposta e autorizará o início dos serviços.")

doc.add_heading("8 Condições de aceite e homologação", level=1)
doc.add_paragraph(
    "A entrega será considerada realizada quando a versão com o módulo de inventário de rua for disponibilizada para testes. A contratante deverá validar o fluxo, os arquivos importados, os resultados do confronto e os relatórios. Ajustes identificados nessa etapa serão atendidos quando estiverem diretamente relacionados ao escopo contratado."
)

doc.add_heading("9 Limites do escopo", level=1)
doc.add_paragraph(
    "Esta proposta não contempla alterações no layout do picking, integrações com sistemas de terceiros, hospedagem, aquisição de licenças, mudanças posteriores nos arquivos de origem ou novas funcionalidades não descritas neste documento. Solicitações adicionais serão analisadas e, quando necessário, apresentadas em novo orçamento."
)

doc.add_heading("10 Validade da proposta", level=1)
doc.add_paragraph("Esta proposta possui validade de 10 dias corridos a partir da data de emissão.")

doc.add_paragraph()
date_p = doc.add_paragraph("Peruíbe SP, 23 de setembro de 2026")
date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

doc.add_paragraph("\n\n____________________________________________")
sig = doc.add_paragraph()
sig.add_run("Hunelyton Mendes Lima").bold = True
doc.add_paragraph("CPF 372.044.388-41")
doc.add_paragraph("E-mail hunelyton.mendes@gmail.com")
doc.add_paragraph("Telefone (13) 98119-6498")

footer = section.footer
footer.distance = Cm(0.7)
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_p.add_run("Hunelyton Mendes Lima   |   Proposta para Drogarias Campeã")
footer_run.font.name = "Aptos"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor(100, 100, 100)

# Prevent short headings from being separated from their first paragraph.
for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith("Heading"):
        keep_with_next(paragraph)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
